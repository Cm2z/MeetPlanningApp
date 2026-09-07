from pathlib import Path
import re
import pdfplumber
from docx import Document

doc = Document(r"P:\MeetPlanning\MeetPlanning_ฉบับแก้ไข_2-Sep.docx")
pdf = r"P:\MeetPlanning\rendered_word\MeetPlanning_ฉบับแก้ไข_2-Sep.pdf"

def norm(s):
    return re.sub(r"\s+", "", s).replace("–", "-").lower()

with pdfplumber.open(pdf) as f:
    pages = [norm(p.extract_text() or "") for p in f.pages]

def find_page(label):
    n = norm(label)
    candidates = [(i+1, len(n), n in txt) for i, txt in enumerate(pages) if i >= 10]
    exact = [i for i, _, ok in candidates if ok]
    return exact[-1] if exact else None

print("MAIN TOC")
for i in range(112, 153):
    if i >= len(doc.paragraphs): break
    text = " ".join(doc.paragraphs[i].text.split())
    if not text or text.startswith("สารบัญ") or text.startswith("เรื่อง"):
        continue
    label = re.sub(r"\s+\S+$", "", text) if re.search(r"\s+(?:\d+|[ก-ฮ](?:-[ก-ฮ])?)$", text) else text
    print(i, repr(text), "=>", find_page(label))

print("TABLE TOC")
for i in range(156, 167):
    text = " ".join(doc.paragraphs[i].text.split())
    print(i, repr(text), "=>", find_page(text))

print("FIGURE TOC")
for i in range(172, 211):
    text = " ".join(doc.paragraphs[i].text.split())
    if text:
        print(i, repr(text), "=>", find_page(text))
