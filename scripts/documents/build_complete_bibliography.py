from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn

SOURCE = r"P:\MeetPlanning\แบบฟอร์มบทที่บรรณานุกรม.docx"
OUTPUT = r"P:\MeetPlanning\บรรณานุกรม_ฉบับครบตามไฟล์ต้นฉบับ.docx"

thai = [
    "รุ่งชีวา และอินทร์จันทร์ (2568) การออกแบบเว็บไซต์ที่ดีต้องคำนึงถึงความต้องการของผู้ใช้งานเป็นหลัก โดยเน้นการใช้งานที่ง่าย รวดเร็ว และปลอดภัย มีระบบนำทางที่ชัดเจน รองรับการแสดงผลบนทุกอุปกรณ์ เข้าถึงได้จาก https://so04.tci-thaijo.org/index.php/SSRUJPD/article/view/279754",
    "มัธยม อ่อนจันทร์, บังอร พลมิตร และศรารัตน์ วรรณแจ่ม (2567) การออกแบบ UI/UX ที่ดีนั้นต้องคำนึงถึงกระบวนการรับรู้ของผู้ใช้งานเป็นหลัก โดยรูปแบบ Flat Design ที่เน้นความเรียบง่าย การใช้สีสดใส และไอคอนแบบแบน ได้รับการประเมินด้านความชัดเจนและความสะดวกในการใช้งานสูงกว่ารูปแบบ Skeuomorphism เข้าถึงได้จาก https://li01.tci-thaijo.org/index.php/sci_01/article/view/264587",
    "ธีระเดชอุปถัมภ์ และคณะ (2568) การออกแบบและพัฒนาเว็บไซต์ที่ดีควรใช้หลักการพื้นฐาน 3 ส่วน ได้แก่ ส่วนหัวของหน้าเว็บไซต์ ส่วนของเนื้อหา และส่วนท้ายของหน้าเว็บไซต์ โดยต้องคำนึงถึงการเลือกใช้สี รูปแบบตัวอักษรที่อ่านง่าย และการรองรับการแสดงผลบนทุกอุปกรณ์ในรูปแบบ Responsive Web Design เข้าถึงได้จาก https://so06.tci-thaijo.org/index.php/tla_bulletin/article/view/288685",
    "ญาณพิทักษ์ และคณะ (2567) อิทธิพลของปัจจัยการตลาดออนไลน์ต่อการตัดสินใจซื้อสินค้าออนไลน์ประกอบด้วยปัจจัยหลัก 4 ด้าน ได้แก่ ความไว้วางใจ การรับรู้ด้านราคา การรับรู้ด้านประโยชน์ และคุณภาพการให้บริการ เข้าถึงได้จาก https://so04.tci-thaijo.org/index.php/WTURJ/article/view/276254",
    "สุวรรณศิลป์ และคณะ (2566) ได้ศึกษาการพัฒนาระบบซื้อขายสินค้าออนไลน์สำหรับร้านค้า โดยออกแบบระบบบริหารจัดการร้านค้าบนเว็บแอปพลิเคชันที่รองรับการสั่งซื้อสินค้า การจัดการฐานข้อมูลสินค้า และการค้นหาข้อมูลได้อย่างรวดเร็ว ผลการประเมินความพึงพอใจของผู้ใช้งานโดยรวมอยู่ในระดับมากที่สุด เข้าถึงได้จาก https://so08.tci-thaijo.org/index.php/romyoongthong/article/view/1993",
    "ลืมสายธาร และศิริสวัสดิ์ (2568) ได้พัฒนาระบบพาณิชย์อิเล็กทรอนิกส์บนแพลตฟอร์มเว็บสำหรับร้านค้าเกษตร เพื่อเพิ่มประสิทธิภาพการบริหารจัดการสินค้าและกระบวนการซื้อขายออนไลน์ ผลการศึกษาพบว่าระบบที่พัฒนาขึ้นสามารถรองรับการจัดการข้อมูลสินค้า การรับคำสั่งซื้อ และการรายงานผลสรุปได้อย่างครบถ้วน เข้าถึงได้จาก https://doi.org/10.57260/stc.2025.1031",
    "เดิมราช และคณะ (2569) ได้พัฒนาระบบสารสนเทศสำหรับการขายบนแพลตฟอร์มเว็บโดยใช้แนวคิดระบบสารสนเทศเพื่อการจัดการเป็นกรอบการพัฒนาร่วมกับวงจรการพัฒนาระบบ ระบบดังกล่าวรองรับการจัดการข้อมูลสินค้า การบันทึกคำสั่งซื้อ และการออกรายงานสรุปผล ผลการประเมินพบว่าระบบมีประสิทธิภาพอยู่ในระดับดีมาก เข้าถึงได้จาก https://doi.org/10.57260/stc.2026.1279",
    "สุบิน ยุระรัช (2565) ทำไมต้องลิเคิร์ต? วารสารนวัตกรรมและการจัดการ ปีที่ 7 ฉบับที่ 1 หน้า 152–165 เข้าถึงได้จาก https://so03.tci-thaijo.org/index.php/journalcim/article/view/259482",
    "รุ่งเรือง มุศิริ, เกริกชัย ดีคำ, ประดิษฐ์ อายุวงษ์, ปวีณา ปรีชญากุล และบุญยฤทธิ์ ศรีปาน (2568) ระบบการจองห้องประชุมออนไลน์ กรณีศึกษา: สำนักงานปลัดกระทรวงกลาโหม (ศรีสมาน) วารสารวิชาการการจัดการเทคโนโลยี มหาวิทยาลัยราชภัฏมหาสารคาม ปีที่ 12 ฉบับที่ 1 หน้า 25–38 เข้าถึงได้จาก https://ph02.tci-thaijo.org/index.php/itm-journal/article/view/256614",
]

english = [
    "Patil et al. (2025) developed a comprehensive e-commerce platform using the MERN Stack — MongoDB, Express.js, React.js, and Node.js. The system supported essential features including a shopping cart, wishlist, online payment gateway, and an administrator management panel. Retrieved from: https://doi.org/10.62110/sciencein.jist.2025.v13.1116",
    "Castillo et al. (2024) investigated the role of online reviews and rating systems in shaping consumer trust on e-commerce platforms, using Mercado Libre Colombia as a case study. Retrieved from: https://doi.org/10.3389/fcomm.2024.1460321",
    "Ma and Wang (2024) proposed a Big Data-assisted E-Commerce Business Model (BD-ECBM) to address security and performance challenges in e-commerce platforms. Retrieved from: https://doi.org/10.1016/j.heliyon.2024.e28571",
]

websites = [
    "World Wide Web Consortium (W3C) (2024) Accessibility, Usability, and Inclusion เข้าถึงได้จาก https://www.w3.org/WAI/fundamentals/accessibility-usability-inclusion/",
    "Level Access (ม.ป.ป.) Usability vs. Accessibility เข้าถึงได้จาก https://www.levelaccess.com/blog/usability-vs-accessibility/",
    "Clinical Informatics Wiki (ม.ป.ป.) Usability เข้าถึงได้จาก http://clinfowiki.org/index.php/Usability",
    "Nielsen Norman Group (ม.ป.ป.) The 3-Click Rule for Navigation Is False เข้าถึงได้จาก https://www.nngroup.com/articles/3-click-rule/",
    "Wikipedia (ม.ป.ป.) Three-click Rule เข้าถึงได้จาก https://en.wikipedia.org/wiki/Three-click_rule",
    "Qualaroo (ม.ป.ป.) 3-Click Rule เข้าถึงได้จาก https://qualaroo.com/blog/3-click-rule/",
    "Nielsen Norman Group (ม.ป.ป.) F-Shaped Pattern of Reading on the Web เข้าถึงได้จาก https://www.nngroup.com/articles/f-shaped-pattern-reading-web-content/",
    "Google Material Design (2024) Material Design 3 เข้าถึงได้จาก https://m3.material.io/",
    "Nielsen Norman Group (2025) Usability Guidelines for Accessible Web Design เข้าถึงได้จาก https://www.nngroup.com/reports/usability-guidelines-accessible-web-design/",
    "MDN Web Docs (2024) Fundamental Text and Font Styling เข้าถึงได้จาก https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Text_styling/Fundamentals",
    "Google UX Research (2025) The Designer’s Guide to Accessibility Research เข้าถึงได้จาก https://design.google/library/designers-guide-accessibility-research",
    "International Organization for Standardization (2019) ISO 9241-210:2019 Ergonomics of Human-System Interaction—Part 210: Human-Centred Design for Interactive Systems เข้าถึงได้จาก https://www.iso.org/standard/77520.html",
    "Microsoft (2025) Visual Studio Code Documentation เข้าถึงได้จาก https://code.visualstudio.com/docs",
    "Node.js (2025) About Node.js เข้าถึงได้จาก https://nodejs.org/en/about/",
    "Express.js (2025) Express Web Framework เข้าถึงได้จาก https://expressjs.com/",
    "Vue.js (2025) Introduction to Vue.js เข้าถึงได้จาก https://vuejs.org/guide/introduction",
    "Oracle Corporation (2025) MySQL เข้าถึงได้จาก https://www.mysql.com/",
    "Apache Friends (2025) XAMPP เข้าถึงได้จาก https://www.apachefriends.org/",
    "MDN Web Docs (2025) HTML: HyperText Markup Language เข้าถึงได้จาก https://developer.mozilla.org/en-US/docs/Web/HTML",
]

doc = Document(SOURCE)
title_tpl = deepcopy(doc.paragraphs[0]._element)
section_tpl = deepcopy(doc.paragraphs[1]._element)
body_tpl = deepcopy(doc.paragraphs[2]._element)
blank_tpl = deepcopy(doc.paragraphs[3]._element)
body = doc._element.body
sect_pr = body.find(qn("w:sectPr"))
for child in list(body):
    if child is not sect_pr:
        body.remove(child)

def append(template, text=""):
    e = deepcopy(template)
    rpr = e.find(".//" + qn("w:rPr"))
    rpr = deepcopy(rpr) if rpr is not None else None
    for child in list(e):
        if child.tag != qn("w:pPr"):
            e.remove(child)
    body.insert(len(body) - 1, e)
    p = next(x for x in doc.paragraphs if x._element is e)
    run = p.add_run(text)
    if rpr is not None:
        run._element.insert(0, rpr)

append(title_tpl, "บรรณานุกรม")
for heading, entries in (("ภาษาไทย", thai), ("ภาษาอังกฤษ", english), ("เว็บไซต์", websites)):
    append(section_tpl, heading)
    for entry in entries:
        append(body_tpl, entry)
        append(blank_tpl)

doc.save(OUTPUT)
print(OUTPUT)
