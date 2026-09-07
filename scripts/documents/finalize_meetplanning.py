from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

PATH = Path(r"P:\MeetPlanning\MeetPlanning_ฉบับแก้ไข_2-Sep.docx")
doc = Document(PATH)


def set_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def clear_direct_numbering(p):
    pPr = p._p.get_or_add_pPr()
    numPr = pPr.find(qn("w:numPr"))
    if numPr is not None:
        pPr.remove(numPr)


# Remove old automatic list numbering from the manually corrected 3.3 hierarchy.
for p in doc.paragraphs:
    if p.text.strip().startswith(("3.3.1 โครงงานซอฟต์แวร์", "3.3.1.1 ", "3.3.1.2 ", "3.3.1.3 ", "3.3.1.4 ", "3.3.1.5 ", "3.3.1.6 ")):
        clear_direct_numbering(p)

# Main contents: replace stale numbers with pages verified from the Word PDF.
main = {
    "บทคัดย่อภาษาไทย": "ก", "บทคัดย่อภาษาอังกฤษ": "ข", "กิตติกรรมประกาศ": "ค",
    "สารบัญ": "ง-จ", "สารบัญตาราง": "ฉ", "สารบัญภาพ": "ช-ซ",
    "บทที่ 1 บทนำ": "1", "1.1 ที่มาและความสำคัญของโครงการ": "1",
    "1.2 วัตถุประสงค์ของโครงการ": "2", "1.3 ขอบเขตด้านระบบ": "2",
    "1.4 ประโยชน์ของระบบ": "3", "1.5 นิยามศัพท์เฉพาะ": "3",
    "บทที่ 2 แนวคิด ทฤษฎี เอกสารและงานวิจัยที่เกี่ยวข้อง": "5",
    "2.1 ทฤษฎีการออกแบบเว็บไซต์": "5", "2.2 องค์ประกอบของการออกแบบเว็บไซต์": "8",
    "2.3 ความหมายของระบบจองและจัดหาห้องประชุม": "8", "2.4 เครื่องมือที่ใช้พัฒนาระบบ": "9",
    "2.5 งานวิจัยที่เกี่ยวข้อง": "14", "2.6 กรอบแนวคิดในการวิจัย": "14",
    "บทที่ 3 วิธีการดำเนินงาน": "15", "3.1 การศึกษาข้อมูลและรวบรวมความต้องการ": "15",
    "3.2 การวางแผนดำเนินโครงงาน": "16", "3.3 การออกแบบระบบ": "16",
    "3.4 เครื่องมือที่ใช้ในการพัฒนา": "19", "3.5 ขั้นตอนการพัฒนาระบบ": "21",
    "3.6 การทดสอบการทำงาน": "22", "3.7 วิธีการเก็บข้อมูลผลการดำเนินงาน": "22",
    "3.8 บันทึกผลการทดสอบของระบบ": "23", "3.9 ภาพหน้าจอ": "23",
    "3.10 ภาพการทดลอง": "27", "3.11 เวลาตอบสนอง": "28",
    "บทที่ 4 ผลการดำเนินงาน": "29", "4.1 ผลการพัฒนาระบบ": "29",
    "4.2 ผลการทดสอบระบบ": "33", "4.3 ผลการทดลองใช้งาน": "34",
    "4.4 คู่มือการติดตั้งและการใช้งาน": "37",
    "บทที่ 5 สรุปผลและข้อเสนอแนะ": "40", "5.1 สรุปผลการดำเนินโครงงาน": "40",
    "5.2 ปัญหาและอุปสรรค": "41", "5.3 ข้อจำกัดของระบบและข้อเสนอแนะแนวทางในการพัฒนา": "42",
    "บรรณานุกรม": "43", "ภาคผนวก": "48", "ก แบบเสนอโครงงาน": "49",
    "ข แผนการดำเนินงาน": "52", "ประวัติผู้จัดทำ": "54",
}

toc_start = next(i for i,p in enumerate(doc.paragraphs) if p.text.strip() == "สารบัญ")
toc_end = next(i for i,p in enumerate(doc.paragraphs[toc_start+1:], toc_start+1) if p.text.strip() == "สารบัญตาราง")
for p in doc.paragraphs[toc_start+1:toc_end]:
    text = " ".join(p.text.split())
    for label, page in main.items():
        if text == label or text.startswith(label + " "):
            set_text(p, f"{label}\t{page}")
            break

# The contents omitted 3.10/3.11 distinction; repurpose the old 3.10 line and add 3.11.
old = next(p for p in doc.paragraphs[toc_start:toc_end] if p.text.strip().startswith("3.10"))
set_text(old, "3.10 ภาพการทดลอง\t27")
new_p = deepcopy(old._element)
old._element.addnext(new_p)
new_para = next(p for p in doc.paragraphs if p._element is new_p)
set_text(new_para, "3.11 เวลาตอบสนอง\t28")

# List of tables.
table_pages = {"3.1":"19", "3.2":"19", "3.3":"20", "3.4":"20", "3.5":"20", "3.6":"21", "3.7":"21", "3.8":"22", "3.9":"28", "4.1":"33", "4.2":"33", "4.3":"37"}
for p in doc.paragraphs:
    s = " ".join(p.text.split())
    if s.startswith("ตารางที่ "):
        parts = s.split()
        if len(parts) >= 2 and parts[1] in table_pages:
            set_text(p, f"{s}\t{table_pages[parts[1]]}")
    if s == "สารบัญภาพ":
        break

# List of figures.
figure_pages = {
    "2.1":"6", "2.2":"8", "2.3":"10", "2.4":"10", "2.5":"11", "2.6":"12", "2.7":"12", "2.8":"13", "2.9":"14",
    "3.1":"15", "3.2":"17", "3.3":"17", "3.4":"17", "3.5":"18", "3.6":"18", "3.7":"19", "3.8":"23", "3.9":"24",
    "3.10":"24", "3.11":"25", "3.12":"25", "3.13":"26", "3.14":"27", "3.15":"27", "3.16":"28",
    "4.1":"29", "4.2":"30", "4.3":"30", "4.4":"31", "4.5":"31", "4.6":"32", "4.7":"33", "4.8":"35", "4.9":"35", "4.10":"36", "4.11":"36",
}
fig_start = next(i for i,p in enumerate(doc.paragraphs) if p.text.strip() == "สารบัญภาพ")
body_start = next(i for i,p in enumerate(doc.paragraphs[fig_start+1:], fig_start+1) if p.text.strip() == "บทที่ 1")
for p in doc.paragraphs[fig_start+1:body_start]:
    s = " ".join(p.text.split())
    if s.startswith("ภาพที่ "):
        parts = s.split()
        if len(parts) >= 2 and parts[1] in figure_pages:
            set_text(p, f"{s}\t{figure_pages[parts[1]]}")

# Normalize newly inserted bibliography records to the existing bibliography style.
good_th = next(p for p in doc.paragraphs if p.text.strip().startswith("ญาณพิทักษ์"))
good_en = next(p for p in doc.paragraphs if p.text.strip().startswith("Castillo et al."))
new_starts = (
    "รุ่งเรือง มุศิริ, เกริกชัย", "Apache Friends.", "Express.js.", "Google Material Design.",
    "Google UX Research.", "International Organization for Standardization.", "MDN Web Docs.",
    "Microsoft.", "Nielsen Norman Group. (2025)", "Node.js.", "Oracle Corporation.", "Vue.js.",
    "World Wide Web Consortium (W3C). (2024)",
)
for p in doc.paragraphs:
    if p.text.strip().startswith(new_starts):
        sample = good_th if p.text.strip().startswith("รุ่งเรือง") else good_en
        p.style = sample.style
        if sample._p.pPr is not None:
            # Preserve placement but use the bibliography paragraph's spacing/indent settings.
            pPr = deepcopy(sample._p.pPr)
            old_pPr = p._p.pPr
            if old_pPr is not None:
                p._p.remove(old_pPr)
            p._p.insert(0, pPr)
        sample_rPr = sample.runs[0]._r.rPr if sample.runs and sample.runs[0]._r.rPr is not None else None
        for r in p.runs:
            if r._r.rPr is not None:
                r._r.remove(r._r.rPr)
            if sample_rPr is not None:
                r._r.insert(0, deepcopy(sample_rPr))

doc.save(PATH)
print(PATH)
