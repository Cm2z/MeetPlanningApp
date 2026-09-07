from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn


SOURCE = r"P:\MeetPlanning\แบบฟอร์มบทที่บรรณานุกรม.docx"
OUTPUT = r"P:\MeetPlanning\บรรณานุกรม_ตรวจตรงบทที่2.docx"


thai_entries = [
    "มัธยม อ่อนจันทร์, บังอร พลมิตร และศรารัตน์ วรรณแจ่ม (2567) Skeuomorphism and Flat Design: การออกแบบ UX/UI เพื่อให้สอดคล้องกับกระบวนการรับรู้ของมนุษย์ เข้าถึงได้จาก https://li01.tci-thaijo.org/index.php/sci_01/article/view/264587",
    "รุ่งเรือง มุศิริ, เกริกชัย ดีคำ, ประดิษฐ์ อายุวงษ์, ปวีณา ปรีชญากุล และบุญยฤทธิ์ ศรีปาน (2568) ระบบการจองห้องประชุมออนไลน์ กรณีศึกษา: สำนักงานปลัดกระทรวงกลาโหม (ศรีสมาน) วารสารวิชาการการจัดการเทคโนโลยี มหาวิทยาลัยราชภัฏมหาสารคาม ปีที่ 12 ฉบับที่ 1 หน้า 25–38 เข้าถึงได้จาก https://ph02.tci-thaijo.org/index.php/itm-journal/article/view/256614",
]

website_entries = [
    "World Wide Web Consortium (W3C) (2024) Accessibility, Usability, and Inclusion เข้าถึงได้จาก https://www.w3.org/WAI/fundamentals/accessibility-usability-inclusion/",
    "Google Material Design (2024) Material Design 3 เข้าถึงได้จาก https://m3.material.io/",
    "Nielsen Norman Group (2025) Usability Guidelines for Accessible Web Design เข้าถึงได้จาก https://www.nngroup.com/reports/usability-guidelines-accessible-web-design/",
    "MDN Web Docs (2024) Fundamental Text and Font Styling เข้าถึงได้จาก https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Text_styling/Fundamentals",
    "Google UX Research (2025) The Designer’s Guide to Accessibility Research เข้าถึงได้จาก https://design.google/library/designers-guide-accessibility-research",
    "Google (2024) แหล่งที่มาภาพองค์ประกอบของการออกแบบเว็บไซต์ เข้าถึงได้จาก https://images.google.com/",
    "International Organization for Standardization (ISO) (2025) ISO 9241-210 Ergonomics of Human-System Interaction—Part 210: Human-Centred Design for Interactive Systems เข้าถึงได้จาก https://www.iso.org/standard/77520.html",
    "Microsoft (2025) Visual Studio Code Documentation เข้าถึงได้จาก https://code.visualstudio.com/docs",
    "Node.js (2025) About Node.js เข้าถึงได้จาก https://nodejs.org/en/about/",
    "Express.js (2025) Express Web Framework เข้าถึงได้จาก https://expressjs.com/",
    "Vue.js (2025) Introduction to Vue.js เข้าถึงได้จาก https://vuejs.org/guide/introduction",
    "Oracle Corporation (2025) MySQL เข้าถึงได้จาก https://www.mysql.com/",
    "Apache Friends (2025) XAMPP เข้าถึงได้จาก https://www.apachefriends.org/",
    "MDN Web Docs (2025) HTML: HyperText Markup Language เข้าถึงได้จาก https://developer.mozilla.org/en-US/docs/Web/HTML",
]


def set_text(paragraph, text):
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


doc = Document(SOURCE)
title_tpl = deepcopy(doc.paragraphs[0]._element)
section_tpl = deepcopy(doc.paragraphs[1]._element)
body_tpl = deepcopy(doc.paragraphs[2]._element)
blank_tpl = deepcopy(doc.paragraphs[3]._element)

body = doc._element.body
sect_pr = body.find(qn("w:sectPr"))
for child in list(body):
    if child is not sect_pr:
        body.remove(child)


def append_clone(template, text=""):
    element = deepcopy(template)
    # The template contains embedded hyperlink XML. Keep paragraph and run
    # formatting, but remove all old visible content before inserting text.
    first_rpr = element.find(".//" + qn("w:rPr"))
    first_rpr = deepcopy(first_rpr) if first_rpr is not None else None
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    body.insert(len(body) - 1, element)
    paragraph = next(p for p in doc.paragraphs if p._element is element)
    run = paragraph.add_run(text)
    if first_rpr is not None:
        run._element.insert(0, first_rpr)
    return paragraph


append_clone(title_tpl, "บรรณานุกรม")
append_clone(section_tpl, "ภาษาไทย")
for entry in thai_entries:
    append_clone(body_tpl, entry)
    append_clone(blank_tpl)

append_clone(section_tpl, "เว็บไซต์")
for entry in website_entries:
    append_clone(body_tpl, entry)
    append_clone(blank_tpl)

doc.save(OUTPUT)
print(OUTPUT)
