from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

SRC = Path(r"P:\MeetPlanning\MeetPlanning_ฉบับแก้ไข_2-Sep.docx")
OUT = Path(r"P:\MeetPlanning\MeetPlanning_ฉบับแก้ไขล่าสุด_2-Sep.docx")
doc = Document(SRC)

def set_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)

def delete_p(p):
    p._element.getparent().remove(p._element)

# Make Table 3.9 use exactly the same dataset and structure as Table 4.2.
t39 = doc.tables[8]
t42 = doc.tables[10]
t39._tbl.getparent().replace(t39._tbl, deepcopy(t42._tbl))

# Use the agreed consistent dataset in both tables: 10 trials for every item.
for table in (doc.tables[8], doc.tables[10]):
    table.rows[0].cells[2].text = "เวลาตอบสนองเฉลี่ย"
    table.rows[1].cells[1].text = "10"

for p in doc.paragraphs:
    s = p.text.strip()
    if s == "ตารางที่ 3.9 ผลการทดสอบเวลาตอบสนอง":
        set_text(p, "ตารางที่ 3.9 ผลการทดสอบเวลาตอบสนองของระบบ")
    elif "ผู้จัดทำบันทึกระยะเวลาที่ระบบใช้ในการตอบสนอง" in s:
        set_text(p, "ผู้จัดทำได้ทดสอบเวลาตอบสนองของแต่ละฟังก์ชันจำนวน 10 ครั้ง และนำผลที่ได้มาคำนวณหาค่าเฉลี่ย ผลการทดสอบพบว่าเวลาตอบสนองเฉลี่ยอยู่ระหว่าง 0.31–0.38 วินาที และทุกรายการผ่านการทดสอบตามเกณฑ์ที่กำหนด")
    elif "(ISO 9241-210, 2025)" in s:
        set_text(p, p.text.replace("(ISO 9241-210, 2025)", "(ISO 9241-210, 2019)"))

# Remove the obsolete note saying Login was tested only once.
for p in list(doc.paragraphs):
    if p.text.strip().startswith("หมายเหตุ: รายการเข้าสู่ระบบทดสอบ 1 ครั้ง"):
        delete_p(p)

# Remove the Yurarat reference if any remains after deleting former section 2.5.
for p in list(doc.paragraphs):
    if "ยูรารัช สุบิน" in p.text or "สุบิน ยุรารัช" in p.text:
        delete_p(p)

# Remove only exact duplicate website bibliography records now covered by the
# complete author-year entries above; retain distinct 3-click/F-pattern sources.
for p in list(doc.paragraphs):
    s = p.text.strip()
    if s.startswith("World Wide Web Consortium (W3C). (ม.ป.ป.). Accessibility, usability, and inclusion"):
        delete_p(p)

# Alphabetize the English reference records while preserving their formatting.
refs = next(p for p in doc.paragraphs if p.text.strip() == "References")
web = next(p for p in doc.paragraphs if p.text.strip() == "เว็บไซต์")
body = doc._element.body
children = list(body)
start = children.index(refs._element)
end = children.index(web._element)
records = [el for el in children[start+1:end] if el.tag == qn("w:p") and "".join(el.itertext()).strip()]
for el in records:
    body.remove(el)
for el in sorted(records, key=lambda e: "".join(e.itertext()).strip().lower()):
    web._element.addprevious(el)

doc.save(OUT)
print(OUT)
