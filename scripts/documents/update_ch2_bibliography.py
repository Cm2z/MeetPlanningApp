from copy import deepcopy
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


SOURCE = r"P:\MeetPlanning\MeetPlanning_ฉบับแก้ไข_2-Sep.docx"
OUTPUT = r"P:\MeetPlanning\MeetPlanning_แก้บท2และบรรณานุกรม_2-Sep.docx"


def set_text_keep_format(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


doc = Document(SOURCE)

# Correct the publication year in Chapter 2 to match the actual ISO edition.
for paragraph in doc.paragraphs:
    if "(ISO 9241-210, 2025)" in paragraph.text:
        set_text_keep_format(paragraph, paragraph.text.replace("(ISO 9241-210, 2025)", "(ISO 9241-210, 2019)"))

thai_refs = [
    "มัธยม อ่อนจันทร์, บังอร พลมิตร และศรารัตน์ วรรณแจ่ม. (2567). Skeuomorphism and Flat Design: การออกแบบ UX/UI เพื่อให้สอดคล้องกับกระบวนการรับรู้ของมนุษย์. เข้าถึงได้จาก https://li01.tci-thaijo.org/index.php/sci_01/article/view/264587",
    "รุ่งเรือง มุศิริ, เกริกชัย ดีคำ, ประดิษฐ์ อายุวงษ์, ปวีณา ปรีชญากุล และบุญยฤทธิ์ ศรีปาน. (2568). ระบบการจองห้องประชุมออนไลน์ กรณีศึกษา: สำนักงานปลัดกระทรวงกลาโหม (ศรีสมาน). วารสารวิชาการ การจัดการเทคโนโลยี มหาวิทยาลัยราชภัฏมหาสารคาม, 12(1), 25–38. เข้าถึงได้จาก https://ph02.tci-thaijo.org/index.php/itm-journal/article/view/256614",
]

english_refs = [
    "Apache Friends. (2025). XAMPP. Retrieved from https://www.apachefriends.org/",
    "Express.js. (2025). Express web framework. Retrieved from https://expressjs.com/",
    "Google. (2024). Google Images. Retrieved from https://images.google.com/",
    "Google Material Design. (2024). Material Design 3. Retrieved from https://m3.material.io/",
    "Google UX Research. (2025). The designer’s guide to accessibility research. Retrieved from https://design.google/library/designers-guide-accessibility-research",
    "International Organization for Standardization. (2019). ISO 9241-210:2019 Ergonomics of human-system interaction—Part 210: Human-centred design for interactive systems. Retrieved from https://www.iso.org/standard/77520.html",
    "MDN Web Docs. (2024). Fundamental text and font styling. Retrieved from https://developer.mozilla.org/en-US/docs/Learn_web_development/Core/Text_styling/Fundamentals",
    "MDN Web Docs. (2025). HTML: HyperText Markup Language. Retrieved from https://developer.mozilla.org/en-US/docs/Web/HTML",
    "Microsoft. (2025). Visual Studio Code documentation. Retrieved from https://code.visualstudio.com/docs",
    "Nielsen Norman Group. (2025). Usability guidelines for accessible web design. Retrieved from https://www.nngroup.com/reports/usability-guidelines-accessible-web-design/",
    "Node.js. (2025). About Node.js. Retrieved from https://nodejs.org/en/about/",
    "Oracle Corporation. (2025). MySQL. Retrieved from https://www.mysql.com/",
    "Vue.js. (2025). Introduction. Retrieved from https://vuejs.org/guide/introduction",
    "World Wide Web Consortium (W3C). (2024). Accessibility, usability, and inclusion. Retrieved from https://www.w3.org/WAI/fundamentals/accessibility-usability-inclusion/",
]

paras = doc.paragraphs
bib_idx = next(i for i, p in enumerate(paras) if p.text.strip() == "บรรณานุกรม" and i > 600)
appendix_idx = next(i for i, p in enumerate(paras) if p.text.strip() == "ภาคผนวก" and i > bib_idx)
bib_heading = paras[bib_idx]
appendix_heading = paras[appendix_idx]

# Use a bibliography body paragraph as the formatting template.
body_template = paras[bib_idx + 2]
section_template = paras[bib_idx + 1]

for p in list(doc.paragraphs[bib_idx + 1:appendix_idx]):
    p._element.getparent().remove(p._element)

def insert_before(ref_paragraph, template, text, centered=False):
    new_p = deepcopy(template._element)
    ref_paragraph._element.addprevious(new_p)
    p = next(x for x in doc.paragraphs if x._element is new_p)
    set_text_keep_format(p, text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    return p


insert_before(appendix_heading, section_template, "ภาษาไทย", centered=True)
for item in thai_refs:
    insert_before(appendix_heading, body_template, item)
insert_before(appendix_heading, section_template, "ภาษาอังกฤษ", centered=True)
for item in english_refs:
    insert_before(appendix_heading, body_template, item)

# Keep the appendix on its own page after the shortened bibliography.
appendix_heading.paragraph_format.page_break_before = True

doc.save(OUTPUT)
print(OUTPUT)
