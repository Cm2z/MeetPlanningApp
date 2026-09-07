from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PATH = Path(r"P:\MeetPlanning\MeetPlanning_ฉบับแก้ไขล่าสุด_2-Sep.docx")
doc = Document(PATH)

def set_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]: r.text = ""
    else:
        p.add_run(text)

# Use the exact base URLs supplied in อ้างอิง.docx.
updates = {
    "Microsoft. (2025).": "Microsoft. (2025). Visual Studio Code documentation. Retrieved from https://code.visualstudio.com/",
    "Node.js. (2025).": "Node.js. (2025). About Node.js. Retrieved from https://nodejs.org/",
    "Express.js. (2025).": "Express.js. (2025). Express web framework documentation. Retrieved from https://expressjs.com/",
    "Vue.js. (2025).": "Vue.js. (2025). Introduction. Retrieved from https://vuejs.org/",
    "Oracle Corporation. (2025).": "Oracle Corporation. (2025). MySQL database documentation. Retrieved from https://www.mysql.com/",
    "Apache Friends. (2025).": "Apache Friends. (2025). XAMPP. Retrieved from https://www.apachefriends.org/",
    "Google Material Design. (2024).": "Google Material Design. (2024). Material Design 3. Retrieved from https://m3.material.io/",
}
for p in doc.paragraphs:
    for prefix, replacement in updates.items():
        if p.text.strip().startswith(prefix):
            set_text(p, replacement)

# Add the CSS and JavaScript references supplied in the attachment.
web = next(p for p in doc.paragraphs if p.text.strip() == "เว็บไซต์")
sample = next(p for p in doc.paragraphs if p.text.strip().startswith("MDN Web Docs. (2025). HTML"))
existing = "\n".join(p.text for p in doc.paragraphs)
extra = [
    "MDN Web Docs. (2025). CSS: Cascading Style Sheets. Retrieved from https://developer.mozilla.org/en-US/docs/Web/CSS",
    "MDN Web Docs. (2025). JavaScript. Retrieved from https://developer.mozilla.org/en-US/docs/Web/JavaScript",
]
for text in extra:
    if text not in existing:
        el = deepcopy(sample._element)
        p = next(x for x in doc.paragraphs if x._element is el) if False else None
        for child in list(el):
            if child.tag == qn("w:r"):
                el.remove(child)
        r = OxmlElement("w:r")
        if sample.runs and sample.runs[0]._r.rPr is not None:
            r.append(deepcopy(sample.runs[0]._r.rPr))
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        el.append(r)
        web._element.addprevious(el)

# Re-sort the complete English bibliography after additions.
refs = next(p for p in doc.paragraphs if p.text.strip() == "References")
children = list(doc._element.body)
start, end = children.index(refs._element), children.index(web._element)
records = [e for e in children[start+1:end] if e.tag == qn("w:p") and "".join(e.itertext()).strip()]
for e in records: doc._element.body.remove(e)
for e in sorted(records, key=lambda x: "".join(x.itertext()).strip().lower()):
    web._element.addprevious(e)

doc.save(PATH)
print(PATH)
