from docx import Document
from pathlib import Path

src = Path(r"P:\MeetPlanning\MeetPlanning_2-Sep_14.00_original.docx")
doc = Document(src)
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
terms = ["MeetPlanning", "2.5", "2.6", "2.7", "3.3", "3.4", "3.5", "3.6", "3.7", "3.8", "3.9", "3.10", "4.2", "บรรณานุกรม", "W3C", "Material", "Nielsen", "MDN", "ISO", "Microsoft", "Node.js", "Express.js", "Vue.js", "Oracle", "Apache", "UX Research", "ยูรารัช", "รุ่งเรือง"]
for i, p in enumerate(doc.paragraphs):
    text = " ".join(p.text.split())
    if text and (any(t.lower() in text.lower() for t in terms) or 500 <= i <= 680):
        print(f"P{i:04d}\t{p.style.name}\t{text}")
for ti, table in enumerate(doc.tables):
    print(f"\nTABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
    for ri, row in enumerate(table.rows):
        vals = [" ".join(c.text.split()) for c in row.cells]
        print(f"R{ri:02d}\t" + " | ".join(vals))
