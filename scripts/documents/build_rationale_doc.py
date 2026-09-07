from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUTPUT = Path(r"P:\MeetPlanning\หลักการและเหตุผล_MeetPlanning.docx")

TITLE = "หลักการและเหตุผล"
PARAGRAPHS = [
    (
        "ปัจจุบันการประชุม การอบรม สัมมนา และการจัดกิจกรรมสังสรรค์มีบทบาทสำคัญต่อการดำเนินงาน"
        "ของหน่วยงานภาครัฐ ภาคเอกชน สถาบันการศึกษา ตลอดจนบุคคลทั่วไป ส่งผลให้ความต้องการใช้"
        "ห้องประชุมและสถานที่จัดกิจกรรมเพิ่มขึ้นอย่างต่อเนื่อง อย่างไรก็ตาม การค้นหาและจองสถานที่ในรูปแบบเดิม"
        "ยังมีข้อจำกัดหลายประการ เช่น ผู้ใช้ต้องติดต่อผู้ให้บริการหลายแห่งผ่านช่องทางที่แตกต่างกัน ข้อมูลด้านราคา "
        "ความจุ สิ่งอำนวยความสะดวก และช่วงเวลาว่างอาจไม่ครบถ้วนหรือไม่เป็นปัจจุบัน อีกทั้งยังขาดระบบที่ช่วย"
        "เปรียบเทียบสถานที่ตามความเหมาะสม ปัญหาเหล่านี้ทำให้กระบวนการตัดสินใจใช้เวลานาน เกิดความไม่สะดวก "
        "และอาจนำไปสู่ข้อผิดพลาดหรือการจองซ้ำซ้อนได้"
    ),
    (
        "จากปัญหาดังกล่าว คณะผู้จัดทำจึงมีแนวคิดในการพัฒนาเว็บไซต์ MeetPlanning ให้เป็นแพลตฟอร์มกลาง"
        "สำหรับค้นหา เปรียบเทียบ และจองห้องประชุมหรือสถานที่จัดกิจกรรมผ่านระบบออนไลน์ ผู้ใช้งานสามารถตรวจสอบ"
        "รายละเอียดของสถานที่ เช่น ที่ตั้ง ราคา ความจุ รูปภาพ สิ่งอำนวยความสะดวก และตารางเวลาว่าง เพื่อประกอบ"
        "การตัดสินใจก่อนทำรายการจองได้อย่างสะดวก รวดเร็ว และโปร่งใส ขณะเดียวกัน ผู้ให้บริการสามารถเผยแพร่"
        "ข้อมูลสถานที่ จัดการห้องและรายการจอง ตลอดจนติดตามสถานะการให้บริการได้อย่างเป็นระบบ"
    ),
    (
        "การพัฒนาเว็บไซต์ MeetPlanning จึงมีส่วนช่วยลดขั้นตอนและระยะเวลาในการติดต่อประสานงาน เพิ่มความถูกต้อง"
        "ของข้อมูล และยกระดับประสิทธิภาพในการบริหารจัดการการจอง นอกจากนี้ ยังเป็นช่องทางประชาสัมพันธ์ที่ช่วยให้"
        "ผู้ให้บริการเข้าถึงกลุ่มผู้ใช้งานได้กว้างขึ้น ตลอดจนสนับสนุนการนำเทคโนโลยีดิจิทัลมาประยุกต์ใช้ในการบริหาร"
        "จัดการสถานที่ให้มีความทันสมัย สอดคล้องกับความต้องการของผู้ใช้งานในปัจจุบัน และสามารถนำไปต่อยอดเพื่อ"
        "รองรับการใช้งานจริงได้ในอนาคต"
    ),
]


def set_thai_font(run, name="TH Sarabun New", size=Pt(16), bold=False):
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = "TH Sarabun New"
    normal.font.size = Pt(16)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "TH Sarabun New")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    normal._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    set_thai_font(title.add_run(TITLE), size=Pt(18), bold=True)

    for text in PARAGRAPHS:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.widow_control = True
        set_thai_font(p.add_run(text), size=Pt(16))

    core = doc.core_properties
    core.title = TITLE
    core.subject = "หลักการและเหตุผลของโครงการเว็บไซต์ MeetPlanning"
    core.author = "คณะผู้จัดทำโครงการ MeetPlanning"
    core.keywords = "MeetPlanning, ห้องประชุม, ระบบจอง, หลักการและเหตุผล"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
